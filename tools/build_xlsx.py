import csv, re
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = "tools/prospect-report.csv"
OUT = "tools/ShopFlix-prospects.xlsx"

JUNK_EMAIL = re.compile(r"(xxx@|@xxx|example\.|@2x|test@test|your-email|email@)", re.I)

rows = []
with open(SRC, newline="", encoding="utf-8") as f:
    for r in csv.DictReader(f):
        email = (r.get("contactEmail") or "").strip()
        if JUNK_EMAIL.search(email):
            email = ""
        rows.append({
            "domain": r["domain"],
            "score": int(r["riskScore"] or 0),
            "level": r["riskLevel"],
            "products": int(r["products"] or 0),
            "age": r["ageDays"],
            "email": email,
            "issues": r["topIssues"].strip().strip('"'),
        })

# sort: High>Med>Low then by score desc
order = {"High": 0, "Medium": 1, "Low": 2, "?": 3}
rows.sort(key=lambda x: (order.get(x["level"], 9), -x["score"]))

ARIAL = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F2A44")
HEAD_FONT = Font(name=ARIAL, bold=True, color="FFFFFF", size=11)
RISK_FILL = {"High": "FCE4E4", "Medium": "FFF3D6", "Low": "E6F4EA", "?": "EDEDED"}
RISK_FONT = {"High": "B00020", "Medium": "8A6100", "Low": "1E7B34", "?": "666666"}
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ---- Summary sheet ----
s = wb.active
s.title = "Summary"
total = len(rows)
with_email = sum(1 for r in rows if r["email"])
counts = {k: sum(1 for r in rows if r["level"] == k) for k in ["High", "Medium", "Low", "?"]}
reach = sum(1 for r in rows if "unreachable" not in r["issues"].lower())

s["A1"] = "ShopFlix AI — Shopify Prospect List"
s["A1"].font = Font(name=ARIAL, bold=True, size=16, color="1F2A44")
s["A2"] = f"Generated {date.today().isoformat()}  ·  source: 'Powered by Shopify' fingerprint, 24 niches (IN/US/UK)"
s["A2"].font = Font(name=ARIAL, size=10, italic=True, color="666666")

summ = [
    ("Total stores", total),
    ("Reachable / live", reach),
    ("With contact email", with_email),
    ("High risk (best leads)", counts["High"]),
    ("Medium risk", counts["Medium"]),
    ("Low risk", counts["Low"]),
    ("Unreachable / unknown", counts["?"]),
]
r0 = 4
for i, (k, v) in enumerate(summ):
    cell_k = s.cell(row=r0+i, column=1, value=k)
    cell_v = s.cell(row=r0+i, column=2, value=v)
    cell_k.font = Font(name=ARIAL, bold=True, size=11)
    cell_v.font = Font(name=ARIAL, size=11)
    cell_v.alignment = Alignment(horizontal="left")

note_r = r0 + len(summ) + 1
notes = [
    "Risk score = deterministic GMC suspension-risk signals (no/thin refund policy, missing contact email,",
    "fake compare-at pricing, urgency/scarcity, tiny images, thin descriptions, no trust badges).",
    "Higher score = more GMC pain = best fit for a free ShopFlix AI scan.",
    "",
    "NOT yet manually validated. 'Powered by Shopify' fingerprint can include false positives;",
    "confirm each store is live, on Shopify, and mid-segment before outreach. Keep outreach compliant",
    "(personalize, lead with the free scan, clear opt-out, modest volume, prefer US recipients).",
]
for i, t in enumerate(notes):
    c = s.cell(row=note_r+i, column=1, value=t)
    c.font = Font(name=ARIAL, size=10, color="444444")
s.column_dimensions["A"].width = 42
s.column_dimensions["B"].width = 14

# ---- Prospects sheet ----
p = wb.create_sheet("Prospects")
headers = ["#", "Store Domain", "Website", "Risk Score", "Risk Level",
           "Products", "Store Age (days)", "Contact Email", "Top GMC Issues"]
p.append(headers)
for col in range(1, len(headers)+1):
    c = p.cell(row=1, column=col)
    c.fill = HEAD_FILL; c.font = HEAD_FONT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER

for idx, r in enumerate(rows, start=1):
    age = r["age"]
    try:
        age = int(age)
    except (ValueError, TypeError):
        age = None
    row_vals = [idx, r["domain"], f"https://{r['domain']}", r["score"], r["level"],
                r["products"] or None, age, r["email"], r["issues"]]
    p.append(row_vals)
    excel_row = idx + 1
    lvl = r["level"]
    for col in range(1, len(headers)+1):
        c = p.cell(row=excel_row, column=col)
        c.font = Font(name=ARIAL, size=10)
        c.border = BORDER
        c.alignment = Alignment(vertical="center",
                                wrap_text=(col == 9),
                                horizontal="center" if col in (1,4,5,6,7) else "left")
    # link
    link = p.cell(row=excel_row, column=3)
    link.hyperlink = row_vals[2]; link.value = row_vals[2]
    link.font = Font(name=ARIAL, size=10, color="1155CC", underline="single")
    # risk colouring
    lc = p.cell(row=excel_row, column=5)
    lc.fill = PatternFill("solid", fgColor=RISK_FILL.get(lvl, "FFFFFF"))
    lc.font = Font(name=ARIAL, size=10, bold=True, color=RISK_FONT.get(lvl, "000000"))
    p.cell(row=excel_row, column=4).font = Font(name=ARIAL, size=10, bold=True)
    # email styling
    ec = p.cell(row=excel_row, column=8)
    if r["email"]:
        ec.hyperlink = f"mailto:{r['email']}"
        ec.font = Font(name=ARIAL, size=10, color="1155CC", underline="single")
    else:
        ec.value = "—"
        ec.font = Font(name=ARIAL, size=10, color="999999")
        ec.alignment = Alignment(horizontal="center")

widths = {1:5, 2:30, 3:30, 4:11, 5:12, 6:10, 7:15, 8:34, 9:60}
for col, w in widths.items():
    p.column_dimensions[get_column_letter(col)].width = w
p.freeze_panes = "A2"
p.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

wb.save(OUT)
print("wrote", OUT, "rows:", len(rows), "with_email:", with_email)
